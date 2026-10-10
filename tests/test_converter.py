"""変換ツールの単体テスト。"""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from compose_to_docker.interpolate import interpolate
from compose_to_docker.plan import build_plan, render_script
from compose_to_docker.render_xlsx import write_workbook
from compose_to_docker.shellutil import bash_sanity

ROOT = Path(__file__).resolve().parents[1]
MINI = ROOT / "tests" / "fixtures" / "mini" / "compose.yml"
OTEL = Path(r"C:\Users\taka_\Grok\Container_Compose_AWS_OpenTelemetry_Grok\compose.yaml")
BUILD_PUSH = Path(r"C:\Users\taka_\Grok\Container_Compose_Build_Push_v2_from_Codex_Grok\compose.yml")


class InterpolateTest(unittest.TestCase):
    def test_defaults_and_escape(self) -> None:
        self.assertEqual(interpolate("${A:-shop}", {}), "shop")
        self.assertEqual(interpolate("${A:-${B:-shop}}-local", {}), "shop-local")
        self.assertEqual(interpolate("$${MYSQL_ROOT_PASSWORD}", {}), "${MYSQL_ROOT_PASSWORD}")
        self.assertEqual(interpolate("${A:+,${A}}", {}), "")
        self.assertEqual(interpolate("${A:+,${A}}", {"A": "x"}), ",x")
        self.assertEqual(interpolate("${A-default}", {"A": ""}), "")
        self.assertEqual(interpolate("${A:-default}", {"A": ""}), "default")


class MiniComposeTest(unittest.TestCase):
    def test_order_health_and_profile(self) -> None:
        plan = build_plan(MINI, selinux="z")
        names = [svc.name for svc in plan.services]
        self.assertEqual(names, ["db", "web"])
        self.assertNotIn("extra", names)
        script = render_script(plan)
        self.assertEqual(bash_sanity(script), [])
        self.assertIn("mini_lab", script)
        self.assertIn("mini_dbdata", script)
        self.assertIn("--network-alias 'db'", script)
        self.assertIn("--network-alias 'web'", script)
        self.assertIn("--read-only", script)
        self.assertIn("--tmpfs '/tmp:size=16m'", script)
        self.assertIn("--publish '8080:80'", script)
        self.assertIn("mysqladmin ping -p${MYSQL_ROOT_PASSWORD} --silent", script)
        self.assertNotIn("rootpassword", script)
        self.assertIn("'MYSQL_ROOT_PASSWORD=secret'", script)
        self.assertLess(script.index("wait_healthy"), script.index("mini/web:dev"))
        with_profile = build_plan(MINI, profiles={"tools"}, selinux="none")
        self.assertEqual([svc.name for svc in with_profile.services], ["db", "extra", "web"])
        plain = render_script(with_profile)
        self.assertNotIn(":z", plain)

    def test_workbook_font_and_grayscale(self) -> None:
        plan = build_plan(MINI, selinux="z")
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "guide.xlsx"
            write_workbook(plan, path)
            wb = load_workbook(path)
            self.assertIn("表紙", wb.sheetnames)
            self.assertIn("実行順", wb.sheetnames)
            self.assertIn("動作原理と歴史", wb.sheetnames)
            for ws in wb.worksheets:
                self.assertFalse(ws.sheet_view.showGridLines)
                for row in ws.iter_rows():
                    for cell in row:
                        if cell.value is None and not cell.has_style:
                            continue
                        self.assertEqual(cell.font.name, "Meiryo UI", msg=f"{ws.title} {cell.coordinate}")
                        self.assertEqual(cell.font.charset, 128, msg=f"{ws.title} {cell.coordinate}")
                        if cell.font.color is not None and cell.font.color.rgb and cell.font.color.type == "rgb":
                            self.assertTrue(_gray(str(cell.font.color.rgb)), msg=str(cell.font.color.rgb))
                        if cell.fill.fgColor is not None and cell.fill.fgColor.rgb and cell.fill.patternType == "solid":
                            self.assertTrue(_gray(str(cell.fill.fgColor.rgb)), msg=f"{ws.title} {cell.fill.fgColor.rgb}")


class RealComposeTest(unittest.TestCase):
    def test_otel_stack(self) -> None:
        if not OTEL.is_file():
            self.skipTest("otel compose がありません")
        plan = build_plan(OTEL, selinux="z")
        names = [svc.name for svc in plan.services]
        self.assertNotIn("batch-ec2", names)
        self.assertLess(names.index("back"), names.index("front"))
        self.assertLess(names.index("jaeger"), names.index("adot-collector"))
        script = render_script(plan)
        self.assertEqual(bash_sanity(script), [])
        self.assertNotIn("C:\\Users", script)
        self.assertNotIn("\r", script)
        self.assertIn("eap-adot-otel_lab", script)
        self.assertIn("eap-adot-otel_aurora-data", script)
        self.assertIn("--network-alias 'front'", script)
        self.assertIn("--network-alias 'back'", script)
        self.assertIn("--publish '8080:8080'", script)
        self.assertIn("--publish '18080:18080'", script)
        self.assertIn("--read-only", script)
        self.assertIn("--tmpfs '/run/eap:uid=6301,gid=6302,mode=0770'", script)
        self.assertIn("aws.ecs.cluster.name=shopdemo-local,aws.ecs.task.family=intra-api-local", script)
        self.assertIn("mysqladmin ping -h 127.0.0.1 -u root -p${MYSQL_ROOT_PASSWORD} --silent", script)
        self.assertIn("'MYSQL_ROOT_PASSWORD=rootpassword'", script)
        self.assertIn("app-eap-base:1.0", script)
        self.assertIn("base/Containerfile", script)
        self.assertLess(script.index("wait_healthy 'eap-back'"), script.index("--name 'eap-front'"))
        self.assertIn(":ro,z", script)
        self.assertNotIn("docker compose", script.replace("#", ""))

    def test_build_push_secrets(self) -> None:
        if not BUILD_PUSH.is_file():
            self.skipTest("compose.yml がありません")
        plan = build_plan(BUILD_PUSH, selinux="none")
        script = render_script(plan)
        self.assertEqual(bash_sanity(script), [])
        self.assertIn("id=jboss_master_password,env=JBOSS_MASTER_PASSWORD", script)
        self.assertIn("id=cacerts,src=/dev/null", script)
        self.assertIn("j1/base.local", script)


def _gray(rgb: str) -> bool:
    text = rgb[-6:]
    if len(text) != 6:
        return False
    try:
        red = int(text[0:2], 16)
        green = int(text[2:4], 16)
        blue = int(text[4:6], 16)
    except ValueError:
        return False
    return red == green == blue


if __name__ == "__main__":
    unittest.main()
