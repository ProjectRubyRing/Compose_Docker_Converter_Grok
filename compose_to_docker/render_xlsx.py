"""解説を、Meiryo UI のモノトーンな Excel ブックにする。"""

from __future__ import annotations

import math
import unicodedata
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from compose_to_docker.model import Plan
from compose_to_docker.render_md import procedure_blocks

INK = "FF222222"
WHITE = "FFFFFFFF"
HEADER_FILL = "FF2F2F2F"
TITLE_FILL = "FF1A1A1A"
BAND = "FFE6E6E6"
ALT = "FFF4F4F4"
LINE = "FFB5B5B5"
PAPER = "FFFFFFFF"

FONT_NAME = "Meiryo UI"


def write_workbook(plan: Plan, path: Path) -> None:
    wb = Workbook()
    _set_default_font(wb)
    cover = wb.active
    cover.title = "表紙"
    _cover(cover, plan)
    _procedure(wb.create_sheet("実行手順"), plan)
    _steps(wb.create_sheet("実行順"), plan)
    _flags(wb.create_sheet("フラグ詳解"), plan)
    _concepts(wb.create_sheet("動作原理と歴史"), plan)
    _envs(wb.create_sheet("環境変数"), plan)
    _builds(wb.create_sheet("ビルドコンテキスト"), plan)
    _mappings(wb.create_sheet("Compose対応表"), plan)
    _stop(wb.create_sheet("停止"), plan)
    wb.active = 0
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _set_default_font(wb: Workbook) -> None:
    default = wb._fonts[0]
    default.name = FONT_NAME
    default.sz = 11
    default.charset = 128


def _font(size: float = 11, bold: bool = False, color: str = INK) -> Font:
    return Font(name=FONT_NAME, size=size, bold=bold, color=color, charset=128)


def _fill(hex_color: str) -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=hex_color, bgColor=hex_color)


def _border() -> Border:
    side = Side(style="thin", color=LINE)
    return Border(left=side, right=side, top=side, bottom=side)


def _align(horizontal: str = "left") -> Alignment:
    return Alignment(wrap_text=True, vertical="top", horizontal=horizontal)


def _prepare(ws: Worksheet, title: str, headers: list[str], widths: list[float]) -> None:
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 110
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.horizontalCentered = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.4
    ws.page_margins.right = 0.4
    ws.page_margins.top = 0.6
    ws.page_margins.bottom = 0.5
    ws.oddHeader.left.text = f'&"{FONT_NAME}"&10Docker コマンド解説'
    ws.oddFooter.left.text = f'&"{FONT_NAME}"&9RHEL 9.8'
    ws.oddFooter.right.text = f'&"{FONT_NAME}"&9&P / &N'
    ws.sheet_properties.tabColor = "FF4A4A4A"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    cell = ws.cell(1, 1, title)
    cell.font = _font(16, True, WHITE)
    cell.fill = _fill(TITLE_FILL)
    cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left")
    for col in range(1, len(headers) + 1):
        extra = ws.cell(1, col, title if col == 1 else None)
        extra.font = _font(16, True, WHITE)
        extra.fill = _fill(TITLE_FILL)
        extra.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left")
        extra.border = _border()
    ws.row_dimensions[1].height = 32
    for col, header in enumerate(headers, start=1):
        item = ws.cell(2, col, header)
        item.font = _font(11, True, WHITE)
        item.fill = _fill(HEADER_FILL)
        item.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        item.border = _border()
        ws.column_dimensions[get_column_letter(col)].width = widths[col - 1]
    ws.row_dimensions[2].height = 30
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(headers))}2"
    ws.print_title_rows = "1:2"
    ws.sheet_view.view = "normal"
    ws.page_setup.scale = None


def _put_row(ws: Worksheet, row: int, values: list[str], widths: list[float], alt: bool) -> None:
    fill = _fill(ALT if alt else PAPER)
    font = _font(11, False, INK)
    for col, value in enumerate(values, start=1):
        cell = ws.cell(row, col, value)
        cell.font = font
        cell.fill = fill
        cell.alignment = _align("center" if col == 1 and _is_short_number(value) else "left")
        cell.border = _border()
    ws.row_dimensions[row].height = _height(values, widths)


def _is_short_number(value: str) -> bool:
    return value.isdigit() and len(value) <= 4


def _height(values: list[str], widths: list[float]) -> float:
    lines = 1
    for value, width in zip(values, widths):
        text = "" if value is None else str(value)
        count = 0
        for para in text.split("\n") or [""]:
            units = sum(2 if _wide(ch) else 1 for ch in para)
            count += max(1, math.ceil(units / max(width - 1, 8)))
        lines = max(lines, count)
    return min(409, max(22, lines * 16 + 8))


def _wide(ch: str) -> bool:
    return unicodedata.east_asian_width(ch) in {"F", "W"}


def _chunks(text: str, size: int = 18) -> list[str]:
    lines = text.splitlines() or [""]
    return ["\n".join(lines[i : i + size]) for i in range(0, len(lines), size)]


def _cover(ws: Worksheet, plan: Plan) -> None:
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 120
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.5
    ws.page_margins.right = 0.5
    ws.sheet_properties.tabColor = "FF4A4A4A"
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 120
    rows = [
        ("文書名", f"{plan.project_name} を Docker コマンドだけで動かす"),
        ("読む順", "表紙 → 実行手順 → 実行順 → フラグ詳解 → 動作原理と歴史 → 環境変数 → ビルドコンテキスト → Compose対応表 → 停止"),
        ("対象ファイル", plan.compose_path),
        ("プロジェクト名", plan.project_name),
        ("前提", "RHEL 9.8。Docker Engine は導入済み。Podman の互換コマンドではないこと。"),
        ("時点", "2026年10月。Docker Engine 29.9.0（2026-10-08）と RHEL 9.8（2026-05-19、kernel 5.14.0-687.5.1.el9_8）を参照しています。"),
        ("全体像", plan.overview),
        ("たとえ", plan.story),
        ("セルの約束", "フォントは Meiryo UI、文字セットは日本語（charset 128）です。色は黒、白、灰色だけです。文章は折り返し、行の高さは文章の量に合わせています。先頭2行は固定し、2行目で絞り込めます。"),
    ]
    title = ws.cell(1, 1, "Docker コマンド解説")
    title.font = _font(20, True, WHITE)
    title.fill = _fill(TITLE_FILL)
    title.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left")
    ws.merge_cells("A1:B1")
    ws.cell(1, 2).font = _font(20, True, WHITE)
    ws.cell(1, 2).fill = _fill(TITLE_FILL)
    ws.row_dimensions[1].height = 36
    for index, (label, value) in enumerate(rows, start=2):
        left = ws.cell(index, 1, label)
        right = ws.cell(index, 2, value)
        left.font = _font(11, True, INK)
        right.font = _font(11, False, INK)
        left.fill = _fill(BAND)
        right.fill = _fill(PAPER if index % 2 == 0 else ALT)
        left.alignment = _align()
        right.alignment = _align()
        left.border = _border()
        right.border = _border()
        ws.row_dimensions[index].height = _height([label, value], [28, 120])
    ws.freeze_panes = "A2"
    ws.print_title_rows = "1:1"
    ws.oddFooter.right.text = f'&"{FONT_NAME}"&9&P / &N'


def _procedure(ws: Worksheet, plan: Plan) -> None:
    headers = ["番号", "項目", "内容"]
    widths = [10, 32, 110]
    _prepare(ws, "RHEL 9.8 で実行する手順", headers, widths)
    for index, (title, body) in enumerate(procedure_blocks(plan), start=1):
        _put_row(ws, index + 2, [str(index), title, body], widths, index % 2 == 0)
    _finish_filter(ws, len(procedure_blocks(plan)))


def _steps(ws: Worksheet, plan: Plan) -> None:
    headers = ["順番", "段階", "タイトル", "何をする", "実行方法", "成功すると", "失敗すると", "動作原理", "たとえ", "歴史", "なぜ必要か", "コマンド"]
    widths = [8, 14, 28, 42, 32, 28, 28, 32, 28, 32, 28, 72]
    _prepare(ws, "実行順。コマンドは続きの行に分かれています", headers, widths)
    row = 3
    alt = False
    for number, step in enumerate(plan.steps, start=1):
        chunks = _chunks(step.script, 16)
        for part, chunk in enumerate(chunks, start=1):
            values = [
                str(number),
                step.phase,
                step.title if part == 1 else f"{step.title}（続き {part}）",
                step.summary if part == 1 else "上のコマンドの続きです。",
                step.how_to if part == 1 else "",
                step.success if part == 1 else "",
                step.failure if part == 1 else "",
                step.principle if part == 1 else "",
                step.analogy if part == 1 else "",
                step.history if part == 1 else "",
                step.why if part == 1 else "",
                chunk,
            ]
            _put_row(ws, row, values, widths, alt)
            row += 1
        alt = not alt
    _finish_filter(ws, row - 3)


def _flags(ws: Worksheet, plan: Plan) -> None:
    headers = ["順番", "段階", "対象", "フラグ", "値", "意味", "実行すると起きること", "たとえ", "背景"]
    widths = [8, 14, 24, 24, 36, 42, 42, 32, 36]
    _prepare(ws, "フラグの詳解。環境変数の1つずつの意味は『環境変数』シートです", headers, widths)
    row = 3
    alt = False
    for number, step in enumerate(plan.steps, start=1):
        for flag in step.flags:
            _put_row(
                ws,
                row,
                [str(number), step.phase, flag.target, flag.flag, flag.value, flag.meaning, flag.happens, flag.analogy, flag.background],
                widths,
                alt,
            )
            row += 1
            alt = not alt
    _finish_filter(ws, row - 3)


def _concepts(ws: Worksheet, plan: Plan) -> None:
    headers = ["用語", "小学生向けのたとえ", "動作原理", "頭の中の絵", "歴史と背景", "必要になった経緯"]
    widths = [24, 42, 48, 36, 48, 42]
    _prepare(ws, "動作原理と歴史", headers, widths)
    for index, concept in enumerate(plan.concepts, start=1):
        _put_row(
            ws,
            index + 2,
            [concept.title, concept.analogy, concept.principle, concept.picture, concept.history, concept.why],
            widths,
            index % 2 == 0,
        )
    _finish_filter(ws, len(plan.concepts))


def _envs(ws: Worksheet, plan: Plan) -> None:
    headers = ["サービス", "変数名", "値", "どこで決めたか", "意味", "元ファイルの説明"]
    widths = [18, 36, 42, 28, 36, 55]
    _prepare(ws, "環境変数。値は変換時に埋めた結果です", headers, widths)
    for index, item in enumerate(plan.env_rows, start=1):
        note = item.file_note
        if len(note) > 4000:
            note = note[:4000] + "\n（長いのでここで切っています）"
        _put_row(ws, index + 2, [item.service, item.key, item.value, item.source, item.meaning, note], widths, index % 2 == 0)
    _finish_filter(ws, len(plan.env_rows))


def _builds(ws: Worksheet, plan: Plan) -> None:
    headers = ["対象", "ファイル", "ビルドコンテキストから確認したこと"]
    widths = [28, 42, 90]
    _prepare(ws, "ビルドコンテキストの確認", headers, widths)
    for index, item in enumerate(plan.build_rows, start=1):
        _put_row(ws, index + 2, [item.target, item.path, item.fact], widths, index % 2 == 0)
    _finish_filter(ws, len(plan.build_rows))


def _mappings(ws: Worksheet, plan: Plan) -> None:
    headers = ["Compose の項目", "Docker コマンド", "意味"]
    widths = [42, 55, 70]
    _prepare(ws, "Compose と Docker コマンドの対応", headers, widths)
    for index, item in enumerate(plan.mapping_rows, start=1):
        _put_row(ws, index + 2, [item.compose, item.docker, item.meaning], widths, index % 2 == 0)
    _finish_filter(ws, len(plan.mapping_rows))


def _stop(ws: Worksheet, plan: Plan) -> None:
    headers = ["順番", "内容", "コマンド"]
    widths = [10, 70, 80]
    _prepare(ws, "停止。ボリュームとイメージは残します", headers, widths)
    intro = "コンテナを止めて捨て、ネットワークを消します。名前付きボリュームは残るので、データベースの中身は残ります。ファイル全体を実行したときは、起動側の exit 0 でここには入りません。"
    _put_row(ws, 3, ["0", intro, ""], widths, False)
    chunks = _chunks(plan.down_script, 16)
    for index, chunk in enumerate(chunks, start=1):
        _put_row(ws, index + 3, [str(index), "停止コマンドの続きです。" if index > 1 else "この塊を bash で実行すると停止します。", chunk], widths, index % 2 == 0)
    _finish_filter(ws, len(chunks) + 1)


def _finish_filter(ws: Worksheet, data_rows: int) -> None:
    if data_rows <= 0:
        return
    last = 2 + data_rows
    ws.auto_filter.ref = f"A2:{get_column_letter(ws.max_column)}{last}"
